import openpyxl as excel, docx, os, datetime

base_dir = os.path.dirname(__file__)
if not os.path.exists(os.path.join(base_dir, 'docs')):
    os.mkdir(os.path.join(base_dir, 'docs'))
save_dir = os.path.join(base_dir, 'docs')

template_file = os.path.join(base_dir, 'letter.docx')
customer_file = os.path.join(base_dir, 'customer.xlsx')
now = datetime.datetime.now()
date = now.strftime('%Y年%m月%d日')

book = excel.load_workbook(customer_file)
keys = ['2021年10月10日', '●●●御中', '■■■様', '★★★の送付について']
cstyle = {
    '★★★の送付について': True
}
sheet = book.active

cards = []
save_files = []
"""
[
    {'2021年10月10日': date, '●●●御中': '福岡物産', '■■■様': '鈴木次郎', '★★★の送付について': '契約書の送付について},
    {'2021年10月10日': date, '●●●御中': '博多文具店', '■■■様': '田中一郎', '★★★の送付について': '契約書の送付について},
    {'2021年10月10日': date, '●●●御中': '大分水産', '■■■様': '佐藤勝男', '★★★の送付について': '契約書の送付について},
]
"""
for row in sheet.iter_rows():
    values = [c.value for c in row]
    if values[0] is None: break
    card = {}
    data = [date, values[0] + '御中', values[1] + '様', '契約書の送付について']
    for key, d in zip(keys, data):
        card[key] = d
    cards.append(card)
    save_files.append(f"{values[0]}_{values[1]}.docx")
        
# print(cards)
print(save_files)

doc = docx.Document(template_file)

for index, card in enumerate(cards):
    for p in doc.paragraphs:
        for k, v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k, v)
                if k in cstyle:
                    font = p.runs[0].font
                    font.bold = True
                    font.underline = True
                    font.size = docx.shared.Pt(20)
                    
    doc.save(os.path.join(save_dir, save_files[index]))
        